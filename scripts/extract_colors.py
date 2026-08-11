from pathlib import Path
import re

import openpyxl


root = Path(__file__).resolve().parent.parent
workbook = openpyxl.load_workbook(root / "拼豆颜色.xlsx", data_only=True)
colors = []

for sheet_name in workbook.sheetnames:
    sheet = workbook[sheet_name]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        code = row[1] if len(row) > 1 else None
        raw_hex = row[2] if len(row) > 2 else None
        if not code or not raw_hex:
            continue
        # The source sheet contains two visually ambiguous O/0 typos.
        hex_value = str(raw_hex).strip().upper().replace("O", "0")
        if not hex_value.startswith("#"):
            hex_value = f"#{hex_value}"
        if not re.fullmatch(r"#[0-9A-F]{6}", hex_value):
            raise RuntimeError(f"{code} 的 HEX 色值无效：{raw_hex}")
        colors.append(
            {
                "id": str(code).strip(),
                "code": str(code).strip(),
                "hex": hex_value,
                "letter": str(code).strip()[0].upper(),
                "sheet": sheet_name,
            }
        )

if len(colors) != 221 or len({color["code"] for color in colors}) != 221:
    raise RuntimeError("色号表必须包含 221 个唯一色号")

lines = [
    'import type { BeadColor } from "../types";',
    "",
    "export const beadColors: BeadColor[] = [",
]

for color in colors:
    lines.append(
        "  { "
        f'id: "{color["id"]}", code: "{color["code"]}", '
        f'hex: "{color["hex"]}", letter: "{color["letter"]}", '
        f'sheet: "{color["sheet"]}"'
        " },"
    )

lines.extend(["];;".replace(";;", ";"), ""])
(root / "src" / "data").mkdir(parents=True, exist_ok=True)
(root / "src" / "data" / "colors.ts").write_text("\n".join(lines), encoding="utf-8")
